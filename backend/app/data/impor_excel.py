"""Isi tabel penduduk dari file Excel hasil pendataan pengurus."""

from collections import Counter, defaultdict
from datetime import date, datetime
import re
import sys

from openpyxl import load_workbook

from app.core.config import settings
from app.data import db
from app.schemas.penduduk import Alamat, Penduduk

NAMA_SHEET = "Data Penduduk"
BARIS_HEADER = 2  # baris 1 = judul

# Satu-satunya definisi kolom: (field Pydantic, label di Excel, lebar kolom).
# Menjadi acuan tunggal struktur data Excel kependudukan padukuhan.
KOLOM: list[tuple[str, str, int]] = [
    ("kodeKeluarga", "Kode Keluarga", 14),
    ("id", "Kode Warga", 14),
    ("nama", "Nama Lengkap", 24),
    ("jenisKelamin", "Jenis Kelamin", 14),
    ("tempatLahir", "Tempat Lahir", 16),
    ("tanggalLahir", "Tanggal Lahir (yyyy-mm-dd)", 20),
    ("agama", "Agama", 12),
    ("statusPerkawinan", "Status Perkawinan", 16),
    ("pendidikan", "Pendidikan Terakhir", 14),
    ("pekerjaan", "Pekerjaan", 20),
    ("golonganDarah", "Gol. Darah", 10),
    ("statusHubunganKeluarga", "Status dalam KK", 18),
    ("kewarganegaraan", "Kewarganegaraan", 14),
    ("jabatan", "Jabatan", 12),
    ("jalan", "Alamat Jalan", 26),
    ("rt", "RT", 6),
    ("rw", "RW", 6),
    ("desa", "Desa/Kelurahan", 16),
    ("kecamatan", "Kecamatan", 14),
    ("kabupaten", "Kabupaten", 14),
    ("provinsi", "Provinsi", 14),
    ("kodePos", "Kode Pos", 10),
    ("statusKependudukan", "Status Kependudukan", 18),
]

# Kolom yang isinya terkunci ke sedikit pilihan. Dipakai pembangkit template
# untuk memasang dropdown; salah ketik ditangkap Pydantic saat impor.
PILIHAN: dict[str, list[str]] = {
    "jenisKelamin": ["LAKI_LAKI", "PEREMPUAN"],
    "agama": ["ISLAM", "KRISTEN", "KATOLIK", "HINDU", "BUDDHA", "KONGHUCU", "LAINNYA"],
    "statusPerkawinan": ["BELUM_KAWIN", "KAWIN", "CERAI_HIDUP", "CERAI_MATI"],
    "pendidikan": ["TIDAK_BELUM_SEKOLAH", "BELUM_TAMAT_SD", "SD", "SMP", "SMA", "D2", "D3", "D4", "S1", "S2", "S3"],
    "golonganDarah": ["A", "B", "AB", "O", "TIDAK_TAHU"],
    "statusHubunganKeluarga": [
        "KEPALA_KELUARGA", "ISTRI", "ANAK", "FAMILI_LAIN", "LAINNYA",
    ],
    "jabatan": ["WARGA", "DUKUH", "RW", "RT"],
    "statusKependudukan": ["AKTIF", "PINDAH", "MENINGGAL"],
}

_ALAMAT = {
    "jalan", "rt", "rw", "desa", "kecamatan", "kabupaten", "provinsi", "kodePos",
}


def _samakan(teks: object) -> str:
    """Header dibandingkan tanpa peduli huruf besar-kecil & spasi berlebih."""
    return " ".join(str(teks or "").split()).lower()


def petakan_kolom(baris_header: tuple) -> dict[str, int]:
    """Nama field -> indeks kolom, dicocokkan dari label di baris header."""
    ada = {_samakan(v): i for i, v in enumerate(baris_header) if v is not None}
    alias = {
        "status": "status kependudukan",
        "status warga": "status kependudukan",
        "keterangan": "status kependudukan",
        "pindah": "status kependudukan",
    }
    for k, v in alias.items():
        if k in ada and v not in ada:
            ada[v] = ada[k]

    peta, hilang = {}, []
    for field, label, _ in KOLOM:
        i = ada.get(_samakan(label))
        if i is None:
            # statusKependudukan opsional: jika header tidak ada, cari kolom ke-23 atau fallback -1 (default AKTIF)
            if field == "statusKependudukan":
                if len(baris_header) >= 23:
                    peta[field] = 22
                else:
                    peta[field] = -1
                continue
            hilang.append(label)
        else:
            peta[field] = i
    if hilang:
        sys.exit(
            "Kolom berikut tidak ketemu di baris header Excel:\n  "
            + "\n  ".join(hilang)
            + "\n\nJudul kolom tidak boleh diubah — urutannya boleh."
        )
    return peta


def baris_ke_penduduk(nilai: dict[str, str]) -> Penduduk:
    inti = {k: v for k, v in nilai.items() if k not in _ALAMAT}
    alamat = {k: v for k, v in nilai.items() if k in _ALAMAT}
    # `id` = kolom "Kode Warga", kunci yang dijaga pengurus. Bukan dibangkitkan
    # acak: jabatan pengurus menunjuk ke warga tertentu, dan impor menimpa
    # seluruh tabel — id acak akan memutus tautan itu tiap kali impor.
    return Penduduk(alamat=Alamat(**alamat), **inti)  # type: ignore[arg-type]


def _bangun_peta_rt_ke_rw(baris_cells: list, peta: dict[str, int]) -> dict[str, str]:
    """Bangun pemetaan RT -> RW secara dinamis dari warga tetap yang ada di spreadsheet."""
    counts: dict[str, Counter] = defaultdict(Counter)
    idx_rt = peta.get("rt", 15)
    idx_rw = peta.get("rw", 16)
    idx_status = peta.get("statusKependudukan", 22)
    for r in baris_cells[1:]:
        r_vals = [c.value for c in r]
        c_status = str(r_vals[idx_status] if len(r_vals) > idx_status else "").strip().lower()
        if "ngontrak" not in c_status and "sementara" not in c_status:
            rt = str(r_vals[idx_rt] if len(r_vals) > idx_rt else "").strip()
            rw = str(r_vals[idx_rw] if len(r_vals) > idx_rw else "").strip()
            if rt and rw and rt != "None" and rw != "None":
                rt_clean = re.sub(r"^0+", "", rt) or "0"
                rw_clean = re.sub(r"^0+", "", rw) or "0" if rw.isdigit() else rw
                counts[rt_clean][rw_clean] += 1
    hasil: dict[str, str] = {}
    for rt, rw_counter in counts.items():
        if rw_counter:
            hasil[rt] = rw_counter.most_common(1)[0][0]
    return hasil


def baca_xlsx(path: str) -> list[Penduduk]:
    ws = load_workbook(path, data_only=True)[NAMA_SHEET]
    baris_cells = list(ws.iter_rows(min_row=BARIS_HEADER, values_only=False))
    peta = petakan_kolom(tuple(c.value for c in baris_cells[0]))

    # Deteksi pemetaan RT -> RW secara dinamis dari data warga
    peta_rt_rw = _bangun_peta_rt_ke_rw(baris_cells, peta)

    # Deteksi kolom bansos secara dinamis dari label baris header
    kolom_header_vals = [c.value for c in baris_cells[0]]
    kolom_bansos = [
        i for i, v in enumerate(kolom_header_vals) if v and "bansos" in _samakan(v)
    ]
    if not kolom_bansos:
        # Fallback memeriksa kolom setelah kolom kependudukan jika ada
        idx_batas = max(peta.values()) + 1 if peta else 22
        kolom_bansos = list(range(idx_batas, len(kolom_header_vals)))

    idx_status_kependudukan = peta.get("statusKependudukan", 22)

    daftar: list[Penduduk] = []
    kosong: list[int] = []
    baris_ke_nomor: dict[str, list[int]] = {}
    # +1 karena `baris_cells` dimulai dari baris header, dan Excel menghitung dari 1.
    for nomor, r in enumerate(baris_cells[1:], start=BARIS_HEADER + 1):
        r_vals = [c.value for c in r]
        if not r_vals[peta["nama"]] and not r_vals[peta["id"]]:
            continue  # dua-duanya kosong = baris belum diisi, lewati
        nilai = {}
        for field, i in peta.items():
            if i == -1 or i >= len(r_vals):
                val = None
            else:
                val = r_vals[i]
            if val is None:
                nilai[field] = ""
            elif isinstance(val, (datetime, date)):
                nilai[field] = val.strftime("%Y-%m-%d")
            else:
                s = str(val).strip()
                if field == "tanggalLahir" and (" " in s or "T" in s):
                    s = s.split(" ")[0].split("T")[0]
                elif field in ("rt", "rw") and s.isdigit():
                    s = re.sub(r"^0+", "", s) or "0"
                nilai[field] = s

        # Keterangan status tambahan (Pindah, Meninggal, Cerai, RT X (Ngontrak))
        c_status_val = (
            r_vals[idx_status_kependudukan]
            if idx_status_kependudukan >= 0 and len(r_vals) > idx_status_kependudukan
            else None
        )
        c_status_str = str(c_status_val or "").strip()

        status_kependudukan = "AKTIF"
        status_domisili = "TETAP"
        catatan_perkawinan = None
        catatan_kematian = None
        alamat_asal = None

        if "pindah" in c_status_str.lower():
            status_kependudukan = "PINDAH"
        elif "meninggal" in c_status_str.lower():
            status_kependudukan = "MENINGGAL"
            sisa = re.sub(r"(?i)meninggal\s*", "", c_status_str).strip(" :-,()")
            if sisa:
                catatan_kematian = sisa
        elif "cerai" in c_status_str.lower():
            status_kependudukan = "AKTIF"
            nilai["statusPerkawinan"] = "CERAI_HIDUP"
            catatan_perkawinan = "Cerai (Belum Update KK)"
        elif "ngontrak" in c_status_str.lower() or "sementara" in c_status_str.lower():
            status_kependudukan = "AKTIF"
            status_domisili = "KONTRAK"
            m = re.search(r"RT\s*(\d+)", c_status_str, re.IGNORECASE)
            kontrak_rt = re.sub(r"^0+", "", m.group(1)) or "1" if m else "1"
            rt_key = kontrak_rt
            kontrak_rw = (
                peta_rt_rw.get(rt_key)
                or peta_rt_rw.get(kontrak_rt)
                or "19"
            )
            alamat_asal = (
                f"{nilai.get('jalan', '')}, RT {nilai.get('rt', '')}/RW {nilai.get('rw', '')}, "
                f"{nilai.get('desa', '')}, {nilai.get('kecamatan', '')}, {nilai.get('kabupaten', '')}"
            )
            nilai["rt"] = kontrak_rt
            nilai["rw"] = kontrak_rw

        nilai["statusKependudukan"] = status_kependudukan
        nilai["statusDomisili"] = status_domisili
        nilai["catatanPerkawinan"] = catatan_perkawinan
        nilai["catatanKematian"] = catatan_kematian
        nilai["alamatAsal"] = alamat_asal

        # Pendidikan & deteksi sel kuning D2
        pend = (nilai.get("pendidikan") or "").strip().upper()
        cell_pend = r[peta["pendidikan"]]
        pend_fill = (
            cell_pend.fill.start_color.rgb
            if (cell_pend.fill and cell_pend.fill.start_color)
            else None
        )
        # Bedakan highlight kuning spesifik sel pendidikan (D2) dari baris yang diwarnai penuh
        sel_kuning_baris = sum(
            1 for c in r if c.fill and c.fill.start_color and c.fill.start_color.rgb == "FFFFFF00"
        )
        if pend_fill == "FFFFFF00" and sel_kuning_baris <= 2:
            nilai["pendidikan"] = "D2"
        elif pend in ("TIDAK_SEKOLAH", "BELUM_SEKOLAH"):
            nilai["pendidikan"] = "TIDAK_BELUM_SEKOLAH"
        elif pend in ("D2", "D-2", "D 2", "DII", "D-II", "D.2", "D.II"):
            nilai["pendidikan"] = "D2"
        elif pend in ("D3", "D-3", "D 3", "DIII", "D-III", "D.3", "D.III"):
            nilai["pendidikan"] = "D3"
        elif pend in ("D4", "D-4", "D 4", "DIV", "D-IV", "D.4", "D.IV"):
            nilai["pendidikan"] = "D4"
        elif pend in ("S1", "S-1", "S 1", "SI"):
            nilai["pendidikan"] = "S1"
        elif pend in ("S2", "S-2", "S 2", "SII"):
            nilai["pendidikan"] = "S2"
        elif pend in ("S3", "S-3", "S 3", "SIII"):
            nilai["pendidikan"] = "S3"
        elif pend:
            nilai["pendidikan"] = pend
        else:
            # Fallback untuk baris yang belum terisi di Excel (mis. row 640)
            nilai["pendidikan"] = "SMA"

        if not nilai.get("statusHubunganKeluarga"):
            nilai["statusHubunganKeluarga"] = "KEPALA_KELUARGA"

        if not nilai.get("golonganDarah"):
            nilai["golonganDarah"] = "TIDAK_TAHU"

        # Bansos (pindai seluruh kolom bansos yang terdeteksi)
        bansos = []
        for idx in kolom_bansos:
            if idx < len(r_vals):
                val = str(r_vals[idx] or "").strip().upper()
                if "BPNT" in val and "BPNT" not in bansos:
                    bansos.append("BPNT")
                if "PKH" in val and "PKH" not in bansos:
                    bansos.append("PKH")
        nilai["bansos"] = bansos

        baris_ke_nomor.setdefault(nilai["id"], []).append(nomor)
        if not nilai["id"]:
            kosong.append(nomor)
            continue
        try:
            daftar.append(baris_ke_penduduk(nilai))
        except Exception as e:
            sys.exit(
                f"Data kependudukan tidak valid di baris {nomor} "
                f"(Kode Warga: {nilai.get('id', '-')}, Nama: {nilai.get('nama', '-')}):\n  {e}\n\n"
                "Periksa kembali baris tersebut di Excel, betulkan nilainya, lalu jalankan ulang."
            )

    if kosong:
        sys.exit(
            "Kolom 'Kode Warga' kosong di baris: "
            + ", ".join(str(b) for b in kosong)
            + "\n\nKode Warga adalah identitas warga di sistem — tanpa itu barisnya "
            "tidak bisa dipakai. Isi dulu, lalu jalankan ulang."
        )

    _periksa_jabatan(daftar, baris_ke_nomor)

    ganda = {k: b for k, b in _cari_ganda(daftar, baris_ke_nomor).items()}
    if ganda:
        sys.exit(
            "Kode Warga berikut dipakai lebih dari satu baris:\n  "
            + "\n  ".join(f"{kode} (baris {', '.join(map(str, b))})" for kode, b in ganda.items())
            + "\n\nSatu kode = satu orang. Dua warga bertukar kode berarti dua orang "
            "bertukar jabatan tanpa ada yang menyadarinya. Betulkan dulu."
        )
    return daftar


def _periksa_jabatan(
    daftar: list[Penduduk], nomor_baris: dict[str, list[int]]
) -> None:
    """Satu jabatan satu orang, sudah di file Excel-nya.

    Dua orang bertanda `RT` di RT yang sama berarti file itu sendiri tidak tahu
    siapa ketuanya — dan aplikasi tidak boleh menebak. Diperiksa sebelum satu
    baris pun ditulis.
    """
    pemegang: dict[str, list[str]] = {}
    for p in daftar:
        if p.jabatan == "WARGA":
            continue
        if p.jabatan == "DUKUH":
            jabatan = "Dukuh"
        elif p.jabatan == "RW":
            jabatan = f"Ketua RW {p.alamat.rw}"
        else:
            jabatan = f"Ketua RT {p.alamat.rt} (RW {p.alamat.rw})"
        pemegang.setdefault(jabatan, []).append(p.nama)

    bentrok = {k: v for k, v in pemegang.items() if len(v) > 1}
    if bentrok:
        sys.exit(
            "Satu jabatan ditandai untuk lebih dari satu orang:\n  "
            + "\n  ".join(f"{k}: {', '.join(v)}" for k, v in bentrok.items())
            + "\n\nSatu jabatan dipegang satu orang. Betulkan kolom Jabatan dulu."
        )


def _cari_ganda(
    daftar: list[Penduduk], nomor_baris: dict[str, list[int]]
) -> dict[str, list[int]]:
    """Kode yang muncul lebih dari sekali, beserta nomor barisnya."""
    return {p.id: nomor_baris[p.id] for p in daftar if len(nomor_baris[p.id]) > 1}


def main(path_xlsx: str, timpa: bool = False) -> None:
    daftar = baca_xlsx(path_xlsx)
    if not daftar:
        sys.exit("Tidak ada baris terisi — cek lagi apakah baris contoh sudah dihapus.")

    conn = db.buka(settings.DATABASE_FILE)

    # Sejak Tahap 3b aplikasi adalah sumber kebenaran data warga, bukan Excel.
    # Tanpa kunci ini, satu perintah impor yang dijalankan karena kebiasaan
    # menghapus seluruh hasil pendataan yang dikerjakan pengurus di aplikasi.
    if not db.kosong(conn) and not timpa:
        jumlah = len(db.muat(conn))
        conn.close()
        sys.exit(
            f"Database sudah berisi {jumlah} warga, dan impor MENIMPA semuanya.\n\n"
            "Sejak pengurus bisa mengubah data lewat aplikasi, isi database bisa "
            "lebih baru daripada file Excel — termasuk warga baru dan koreksi "
            "yang tidak ada di file ini.\n\n"
            "Kalau memang mau membuang seluruh isi database dan menggantinya "
            "dengan file ini, ulangi dengan:\n"
            f"  .venv/bin/python -m app.data.impor_excel {path_xlsx} --timpa-semua\n\n"
            "Salin dulu file databasenya sebelum itu."
        )

    lama = db.kosongkan(conn)
    masuk = db.simpan(conn, daftar)
    conn.close()
    if lama:
        print(f"PERHATIAN: {lama} baris lama dihapus dan diganti seluruhnya.")
    print(f"OK: {masuk} warga masuk ke {settings.DATABASE_FILE}")
    print("Langsung kepakai — backend query database tiap request, tidak ada cache.")


if __name__ == "__main__":
    arg = [a for a in sys.argv[1:] if a != "--timpa-semua"]
    if len(arg) != 1:
        sys.exit(
            "Pakai: python -m app.data.impor_excel <path-ke-file.xlsx> [--timpa-semua]"
        )
    main(arg[0], timpa="--timpa-semua" in sys.argv)
